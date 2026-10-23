"""Exercise real local converters; fixtures are diagnostics, not course content."""
import hashlib
import json
import sys
import tempfile
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

from convert_document import LegacyDocumentRequired, convert_document


def main():
    root = Path(sys.argv[1]).resolve()
    root.mkdir(parents=True, exist_ok=True)
    samples = []
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    box.text = "DIAGNOSTIC ONLY #1 원고 분석 시험"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "BEFORE"
    table.cell(0, 1).text = "AFTER"
    deck_path = root / "diagnostic.pptx"
    deck.save(deck_path)
    samples.append((deck_path, ["#1", "원고 분석 시험", "BEFORE", "AFTER"]))
    doc = Document()
    doc.add_heading("DIAGNOSTIC ONLY", 0)
    doc.add_paragraph("원고 전처리 #1")
    doc_path = root / "diagnostic.docx"
    doc.save(doc_path)
    samples.append((doc_path, ["원고 전처리 #1"]))
    pdf_path = root / "diagnostic.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawString(72, 720, "DIAGNOSTIC ONLY - PDF TEXT #1")
    pdf.save()
    samples.append((pdf_path, ["PDF TEXT #1"]))
    checks = []
    for path, expected in samples:
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        result = convert_document(path)
        passed = all(token in result["markdown"] for token in expected)
        unchanged = hashlib.sha256(path.read_bytes()).hexdigest() == before
        checks.append({"format": path.suffix, "expected_text_found": passed,
                       "source_unchanged": unchanged, "source_sha256": before})
        (root / (path.name + ".json")).write_text(json.dumps(result, ensure_ascii=False, indent=2))
        assert passed and unchanged, path
    for suffix in [".hwp", ".hwpx"]:
        path = root / ("routing-check" + suffix)
        path.write_bytes(b"DIAGNOSTIC ROUTING CHECK - NOT A HANGUL FILE")
        try:
            convert_document(path)
        except LegacyDocumentRequired:
            checks.append({"format": suffix, "legacy_route_preserved": True,
                           "real_hangul_parsing_tested": False})
        else:
            raise AssertionError("Hangul input must stay with legacy pipeline")
    payload = {"tests": checks, "status": "PASS", "scope": "diagnostic local preprocessing and routing only"}
    (root / "test-results.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2))
    print(json.dumps(payload, ensure_ascii=False))


if __name__ == "__main__":
    main()
