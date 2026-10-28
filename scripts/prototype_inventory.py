#!/usr/bin/env python3
"""Extract structural observations from a PPTX. Human visual review remains required."""
import argparse
import hashlib
import json
from pathlib import Path
from pptx import Presentation


def sha(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''): h.update(chunk)
    return h.hexdigest()


def inventory(path):
    p = Presentation(path)
    pages = []
    for number, slide in enumerate(p.slides, 1):
        shapes = []
        for s in slide.shapes:
            item = {'type': str(s.shape_type), 'left': s.left, 'top': s.top,
                    'width': s.width, 'height': s.height}
            if s.has_text_frame:
                item['text'] = s.text
                item['fonts'] = sorted({r.font.name for para in s.text_frame.paragraphs
                                        for r in para.runs if r.font.name})
            shapes.append(item)
        pages.append({'page': number, 'shapes': shapes})
    return {'sha256': sha(path), 'canvas_emu': [p.slide_width, p.slide_height],
            'pages': pages, 'visual_review_required': True,
            'limits': 'Top-level objects only; inherited styles, groups, colors and semantic roles require inspection.'}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('pptx'); p.add_argument('--output', required=True)
    args = p.parse_args()
    with Path(args.output).open('x', encoding='utf-8') as f:
        json.dump(inventory(args.pptx), f, ensure_ascii=False)


if __name__ == '__main__': main()
